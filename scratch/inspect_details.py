import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = pptx.Presentation('SamplePPT.pptx')
print("Presentation slide size:", prs.slide_width.inches, prs.slide_height.inches)

for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} Layout: '{slide.slide_layout.name}' ---")
    for s in slide.shapes:
        print(f"  Shape: '{s.name}' ({s.shape_type}) at (left={s.left}, top={s.top}, w={s.width}, h={s.height})")
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    font_info = ""
                    if p.runs:
                        r0 = p.runs[0]
                        font_info = f"[Font: {r0.font.name}, Size: {r0.font.size}, Bold: {r0.font.bold}]"
                    print(f"    P: '{p.text[:60]}' {font_info}")
