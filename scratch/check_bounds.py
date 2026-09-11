import pptx

prs = pptx.Presentation('DAKSH_Service_Learning_Presentation.pptx')
width = prs.slide_width
height = prs.slide_height

print(f"Presentation dimensions: {width.inches} x {height.inches} inches")

issues = []
for i, slide in enumerate(prs.slides):
    s_num = i + 1
    for s in slide.shapes:
        # Check if shape is out of bounds
        if s.left < 0 or s.top < 0 or (s.left + s.width) > width or (s.top + s.height) > height:
            # Note: master background shapes may extend slightly, but slide content should not
            if s.name not in ['bg object 16', 'bg object 17', 'bg object 18', 'bg object 19']:
                issues.append(f"Slide {s_num}: Shape '{s.name}' out of bounds (left={s.left.inches:.2f}, top={s.top.inches:.2f}, right={(s.left+s.width).inches:.2f}, bottom={(s.top+s.height).inches:.2f})")

if not issues:
    print("All shapes are within slide bounds!")
else:
    for iss in issues[:10]:
        print(iss)
