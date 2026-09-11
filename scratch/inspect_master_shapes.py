import pptx

prs = pptx.Presentation('SamplePPT.pptx')
print("Slide Master Shapes:")
for s_idx, shape in enumerate(prs.slide_masters[0].shapes):
    print(f"  Master Shape {s_idx}: name='{shape.name}', type={shape.shape_type}")
    if shape.has_text_frame:
        print(f"    Text: {[p.text for p in shape.text_frame.paragraphs]}")

for l_idx, layout in enumerate(prs.slide_layouts):
    print(f"\nLayout {l_idx}: '{layout.name}'")
    for s_idx, shape in enumerate(layout.shapes):
        print(f"  Layout Shape {s_idx}: name='{shape.name}', type={shape.shape_type}")
        if shape.has_text_frame:
            print(f"    Text: {[p.text for p in shape.text_frame.paragraphs]}")

