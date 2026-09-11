import pptx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('SamplePPT.pptx')

with open('scratch/sample_ppt_dump.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total slides: {len(prs.slides)}\n")
    f.write(f"Slide width: {prs.slide_width.inches} in, height: {prs.slide_height.inches} in\n\n")
    for idx, slide in enumerate(prs.slides):
        f.write(f"=== SLIDE {idx+1} (Layout: {slide.slide_layout.name}) ===\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"  Shape {s_idx}: name='{shape.name}', type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}\n")
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    runs_info = ' | '.join([f"'{r.text}'(font={r.font.name}, sz={r.font.size}, b={r.font.bold}, c={getattr(r.font.color, 'rgb', 'none') if r.font.color else 'none'})" for r in p.runs])
                    f.write(f"    P{p_idx} [align={p.alignment}, level={p.level}]: '{p.text}'\n")
                    if runs_info:
                        f.write(f"      Runs: {runs_info}\n")
            elif shape.has_table:
                f.write(f"    Table: {len(shape.table.rows)} rows, {len(shape.table.columns)} cols\n")
                for r_idx, r in enumerate(shape.table.rows):
                    r_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
                    f.write(f"      Row {r_idx}: {r_txt}\n")
            elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                f.write(f"    Picture: {shape.name}\n")
            f.write("\n")
        f.write("\n" + "="*50 + "\n\n")

print("Dump completed to scratch/sample_ppt_dump.txt")
