import pptx

prs = pptx.Presentation('SamplePPT.pptx')
for idx, layout in enumerate(prs.slide_layouts):
    s = prs.slides.add_slide(layout)
    print(f"\nLayout {idx} ('{layout.name}') has {len(s.shapes)} shapes:")
    for shp in s.shapes:
        print(f"  Shape: '{shp.name}' type={shp.shape_type} at left={shp.left}, top={shp.top}, w={shp.width}, h={shp.height}")
