import docx

doc = docx.Document('DAKSH_Weekly_Reflection_Report.docx')
table = doc.tables[0]

data = [
    # Week 1
    {
        "goals": "Understand what DAKSH needs and how legal researchers study court cases.",
        "methodology": "Discussed project goals with the DAKSH team; studied how legal researchers work and identified the key court information they need for their studies.",
        "contributions": [
            "• Julius: Listed the main needs and goals of legal researchers.",
            "\n• Steven: Explored how the eCourts website is organized.",
            "\n• Mishael: Checked if court data can be collected automatically.",
            "\n• Rahul: Listed the checks needed to ensure data is correct.",
            "\n• Abhishan: Planned how to cleanly store and organize the case data."
        ]
    },
    # Week 2
    {
        "goals": "Learn how the eCourts website is organized and how case details are structured.",
        "methodology": "Explored different pages on the eCourts website, including basic case information, hearing dates, daily updates, court notices, and final judgments.",
        "contributions": [
            "• Julius: Identified the most important case details researchers need.",
            "\n• Steven: Mapped how to find cases across different court levels.",
            "\n• Mishael: Created simple trial tools to gather case information automatically.",
            "\n• Rahul: Categorized different court dates (hearings, notices, decisions).",
            "\n• Abhishan: Designed a clean table layout to store all case details."
        ]
    },
    # Week 3
    {
        "goals": "Build and test an automated process to collect court case information.",
        "methodology": "Created automated tools to visit the eCourts website, search for specific cases, and collect information from every section of the case page.",
        "contributions": [
            "• Julius: Confirmed that all necessary legal information is being collected.",
            "\n• Steven: Improved how the tool navigates through different court pages.",
            "\n• Mishael: Set up the automated data collection pipeline.",
            "\n• Rahul: Checked the collected information for any missing details.",
            "\n• Abhishan: Organized the collected records so they are easy to save and view."
        ]
    },
    # Week 4
    {
        "goals": "Verify the collected data for accuracy and organize all details into proper categories.",
        "methodology": "Cross-checked the gathered data against the official website for accuracy, clearly separated different date types (hearings, notices, orders, case closing), and removed duplicate entries.",
        "contributions": [
            "• Julius: Reviewed the data to make sure it is meaningful and easy to interpret for research.",
            "\n• Steven: Handled multi-page search results and website loading issues.",
            "\n• Mishael: Refined the data collection tools to place details into the correct categories.",
            "\n• Rahul: Double-checked that all recorded court dates match the official records.",
            "\n• Abhishan: Cleaned up the data tables to avoid repetitive information."
        ]
    },
    # Week 5
    {
        "goals": "Plan a user-friendly case search platform and outline future smart features.",
        "methodology": "Designed an easy-to-use search and filtering layout for researchers to browse cases easily, and discussed future additions like automated case summaries and trend reports.",
        "contributions": [
            "• Julius: Outlined useful search filters and categories for researchers.",
            "\n• Steven: Planned how case details and timeline history should appear on screen.",
            "\n• Mishael: Connected the data so researchers can easily view and download it.",
            "\n• Rahul: Created simple quality checklists to ensure data stays reliable.",
            "\n• Abhishan: Designed the layout and screens for the case search portal."
        ]
    }
]

for idx, entry in enumerate(data, start=1):
    row = table.rows[idx]
    
    # Cell 2: Goals and Objectives
    row.cells[2].paragraphs[0].runs[0].text = entry["goals"]
    
    # Cell 3: Methodology
    row.cells[3].paragraphs[0].runs[0].text = entry["methodology"]
    
    # Cell 4: Team Work
    for c_idx, contrib_text in enumerate(entry["contributions"]):
        row.cells[4].paragraphs[0].runs[c_idx].text = contrib_text

doc.save('DAKSH_Weekly_Reflection_Report.docx')
print("Updated DAKSH_Weekly_Reflection_Report.docx successfully.")
