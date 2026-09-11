import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation('SamplePPT.pptx')
print("Loaded SamplePPT.pptx successfully.")
print("Total initial slides:", len(prs.slides))
