import pptx
import os

filepath = os.path.abspath('DAKSH_Service_Learning_Presentation.pptx')
print("Absolute path:", filepath)
prs = pptx.Presentation(filepath)
print(f"Total slides: {len(prs.slides)}")
print(f"Width: {prs.slide_width.inches} in, Height: {prs.slide_height.inches} in")
print(f"Slide width EMU: {prs.slide_width}, Slide height EMU: {prs.slide_height}")

for idx, slide in enumerate(prs.slides):
    print(f"Slide {idx+1}: layout='{slide.slide_layout.name}', shapes={len(slide.shapes)}")
