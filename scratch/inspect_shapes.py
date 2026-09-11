import pptx

prs = pptx.Presentation('SamplePPT.pptx')
print(f"Presentation width={prs.slide_width}, height={prs.slide_height}")

for s_idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {s_idx+1} (Layout: {slide.slide_layout.name}) ---")
    print(f"Background: {slide.background.fill.type if slide.background else 'None'}")
    for shp_idx, shape in enumerate(slide.shapes):
        pos = f"left={shape.left}, top={shape.top}, w={shape.width}, h={shape.height}"
        txt_prev = ""
        if shape.has_text_frame:
            txt_prev = " | ".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()][:3])
            txt_prev = f"Text: [{txt_prev[:80]}]"
        elif shape.has_table:
            txt_prev = f"Table {len(shape.table.rows)}x{len(shape.table.columns)}"
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            txt_prev = f"Picture: {shape.name}"
        else:
            txt_prev = f"Shape: {shape.name} ({shape.shape_type})"
        print(f"  Shape {shp_idx}: {shape.name} ({shape.shape_type}) at {pos} -> {txt_prev}")
