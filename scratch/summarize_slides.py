import pptx

prs = pptx.Presentation('SamplePPT.pptx')
print(f"Total slides in SamplePPT.pptx: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    title = "NO TITLE"
    texts = []
    has_table = False
    has_pic = False
    pic_names = []
    
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
        if s.has_table:
            has_table = True
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            has_pic = True
            pic_names.append(s.name)
            
    title_cand = texts[0] if texts else "EMPTY"
    print(f"\nSlide {i+1}:")
    print(f"  Layout: {slide.slide_layout.name}")
    print(f"  First 3 text lines: {texts[:3]}")
    print(f"  Total text paragraphs: {len(texts)}")
    print(f"  Has Table: {has_table}, Has Picture: {has_pic} ({pic_names})")
    print(f"  Shapes count: {len(slide.shapes)}")
