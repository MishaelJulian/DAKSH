import pandas as pd
import openpyxl
import os

xlsx_path = r"C:\Users\misha\OneDrive\Desktop\daksh\eCourts_Executive_Petitions_2010\Scraped_Cases.xlsx"

if os.path.exists(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    print("Sheets in Scraped_Cases.xlsx:", wb.sheetnames)
    for sheet in wb.sheetnames:
        df = pd.read_excel(xlsx_path, sheet_name=sheet, nrows=2)
        print(f"\nSheet '{sheet}' shape: {df.shape}")
        print("Columns:", list(df.columns))
else:
    print("Scraped_Cases.xlsx does not exist!")
