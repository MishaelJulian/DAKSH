import pptx

prs = pptx.Presentation('SamplePPT.pptx')
print(f"Slide count before: {len(prs.slides)}")

# In python-pptx, we can remove slide elements from prs.slides._sldIdLst
# Let's test removing a slide and adding a slide
rId = prs.slides._sldIdLst[15].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[15]
print(f"Slide count after removing 1 slide: {len(prs.slides)}")

# Add a slide
new_slide = prs.slides.add_slide(prs.slide_layouts[1])
print(f"Slide count after adding 1 slide: {len(prs.slides)}")
