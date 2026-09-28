"""
deeper_tool.exporter
====================
Export pipelines for deeper eCourts data models:
- Multi-sheet styled Excel workbook (.xlsx) with dedicated relational sheets
- Structured JSON output with complete nested case records
- Flat summary CSV output
"""

from __future__ import annotations

import csv
import json
from typing import List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from deeper_tool.db import DeeperDatabase
from deeper_tool.models import DeeperCaseDetail


def export_to_json(db: DeeperDatabase, output_path: str, cnr_list: Optional[List[str]] = None) -> None:
    """Exports cases to a structured JSON file."""
    cnrs = cnr_list or db.list_cnrs()
    cases_data = []
    for cnr in cnrs:
        case = db.get_case(cnr)
        if case:
            cases_data.append(case.to_dict())

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(cases_data, f, indent=2, ensure_ascii=False)


def export_to_csv_flat(db: DeeperDatabase, output_path: str, cnr_list: Optional[List[str]] = None) -> None:
    """Exports cases to a flat CSV file combining summary fields."""
    cnrs = cnr_list or db.list_cnrs()
    headers = [
        "cnr_number", "case_number", "case_type", "court_name",
        "filing_number", "filing_date", "registration_number", "registration_date",
        "decision_date", "case_status", "nature_of_disposal", "court_number_judge",
        "petitioners", "respondents", "acts_sections", "processes_count",
        "main_matters", "hearings_count", "orders_count"
    ]

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)

        for cnr in cnrs:
            case = db.get_case(cnr)
            if not case:
                continue

            pets = "; ".join(f"{p.name} (Adv: {p.advocate})" if p.advocate else p.name for p in case.petitioners)
            resps = "; ".join(f"{r.name} (Adv: {r.advocate})" if r.advocate else r.name for r in case.respondents)
            acts_str = "; ".join(f"{a.act}: {a.section}" if a.section else a.act for a in case.acts)
            mm_str = "; ".join(f"{m.main_case_number} ({m.main_cnr_number})" if m.main_cnr_number else m.main_case_number for m in case.main_matters)

            writer.writerow([
                case.cnr_number, case.case_number, case.case_type, case.court_name,
                case.filing_number, case.filing_date, case.registration_number, case.registration_date,
                case.decision_date, case.case_status, case.nature_of_disposal, case.court_number_judge,
                pets, resps, acts_str, len(case.processes),
                mm_str, len(case.hearings), len(case.orders)
            ])


def export_to_multisheet_excel(db: DeeperDatabase, output_path: str, cnr_list: Optional[List[str]] = None) -> None:
    """Exports full case relational data into a styled multi-sheet Excel workbook."""
    cnrs = cnr_list or db.list_cnrs()
    cases: List[DeeperCaseDetail] = []
    for cnr in cnrs:
        c = db.get_case(cnr)
        if c:
            cases.append(c)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    def style_sheet(ws, col_headers: List[str], data_rows: List[List[Any]]) -> None:
        ws.append(col_headers)
        for r_idx, row in enumerate(data_rows, start=2):
            ws.append(row)

        for col_idx in range(1, len(col_headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for r_idx in range(2, len(data_rows) + 2):
            for col_idx in range(1, len(col_headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = data_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin_border

        ws.row_dimensions[1].height = 28
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)
        ws.freeze_panes = "A2"

    # Sheet 1: Case Overview
    ws_case = wb.create_sheet(title="Case Overview")
    case_headers = [
        "CNR Number", "Case Number", "Case Type", "Court Name",
        "Filing No", "Filing Date", "Registration No", "Registration Date",
        "First Hearing Date", "Decision Date", "Status", "Nature of Disposal",
        "Court Room & Judge"
    ]
    case_rows = [
        [
            c.cnr_number, c.case_number, c.case_type, c.court_name,
            c.filing_number, c.filing_date, c.registration_number, c.registration_date,
            c.first_hearing_date, c.decision_date, c.case_status, c.nature_of_disposal,
            c.court_number_judge
        ]
        for c in cases
    ]
    style_sheet(ws_case, case_headers, case_rows)

    # Sheet 2: Processes
    ws_proc = wb.create_sheet(title="Processes")
    proc_headers = ["Case CNR", "Process ID", "Process Title", "Process Date"]
    proc_rows = [
        [c.cnr_number, pr.process_id, pr.process_title, pr.process_date]
        for c in cases for pr in c.processes
    ]
    style_sheet(ws_proc, proc_headers, proc_rows)

    # Sheet 3: Main Matters
    ws_mm = wb.create_sheet(title="Main Matters")
    mm_headers = ["Case CNR", "Main Case Number", "Main CNR Number", "Main Filing Number"]
    mm_rows = [
        [c.cnr_number, m.main_case_number, m.main_cnr_number, m.main_filing_number]
        for c in cases for m in c.main_matters
    ]
    style_sheet(ws_mm, mm_headers, mm_rows)

    # Sheet 4: Parties & Advocates
    ws_parties = wb.create_sheet(title="Parties & Advocates")
    party_headers = ["Case CNR", "Party Type", "Party Name", "Advocate Name"]
    party_rows = [
        [c.cnr_number, p.type.title(), p.name, p.advocate]
        for c in cases for p in c.parties
    ]
    style_sheet(ws_parties, party_headers, party_rows)

    # Sheet 5: Statutory Acts
    ws_acts = wb.create_sheet(title="Statutory Acts")
    act_headers = ["Case CNR", "Under Act", "Under Section"]
    act_rows = [
        [c.cnr_number, a.act, a.section]
        for c in cases for a in c.acts
    ]
    style_sheet(ws_acts, act_headers, act_rows)

    # Sheet 6: Chronological Hearings
    ws_hearings = wb.create_sheet(title="Hearing History")
    h_headers = ["Case CNR", "Index", "Presiding Judge", "Business Date", "Hearing Date", "Purpose of Hearing"]
    h_rows = [
        [c.cnr_number, h.hearing_index, h.judge, h.business_date, h.hearing_date, h.purpose]
        for c in cases for h in c.hearings
    ]
    style_sheet(ws_hearings, h_headers, h_rows)

    # Sheet 7: Orders & Judgments
    ws_orders = wb.create_sheet(title="Orders & Judgments")
    order_headers = ["Case CNR", "Order Number", "Order Date", "Order Details", "PDF Link Params"]
    order_rows = [
        [c.cnr_number, o.order_number, o.order_date, o.order_details, o.pdf_params]
        for c in cases for o in c.orders
    ]
    style_sheet(ws_orders, order_headers, order_rows)

    wb.save(output_path)
